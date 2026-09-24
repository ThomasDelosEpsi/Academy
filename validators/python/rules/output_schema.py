from __future__ import annotations

import csv
import json
from pathlib import Path

from .base import BaseRule, ValidationContext, infer_scalar_type

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


def _read_output_file(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return [row for row in data if isinstance(row, dict)]
        if isinstance(data, dict):
            return [data]
        return []

    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))

    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"} and load_workbook is not None:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        result = []
        for row in rows[1:]:
            result.append({headers[index]: row[index] for index in range(min(len(headers), len(row))) if headers[index]})
        return result

    return []


class Rule(BaseRule):
    key = "output_schema"
    script_name = "output_schema.py"

    def run(self, context: ValidationContext, rule_config: dict) -> object:
        if context.output_file is None:
            return self.result(
                passed=False,
                severity=rule_config.get("severity", "blocking"),
                message="Aucun fichier de sortie fourni pour verifier le schema.",
            )

        if not context.output_file.exists():
            return self.result(
                passed=False,
                severity=rule_config.get("severity", "blocking"),
                message="Le fichier de sortie attendu n'existe pas.",
                details={"output_file": str(context.output_file)},
            )

        rows = _read_output_file(context.output_file)
        if not rows:
            return self.result(
                passed=False,
                severity=rule_config.get("severity", "blocking"),
                message="Le fichier de sortie est vide ou illisible.",
                details={"output_file": str(context.output_file)},
            )

        headers = set(rows[0].keys())
        expected_columns = context.output_columns or []
        missing = []
        mismatches = []
        for column in expected_columns:
            name = column.get("name")
            if not name:
                continue
            if column.get("required", True) and name not in headers:
                missing.append(name)
                continue
            if name in headers:
                detected_type = infer_scalar_type(rows[0].get(name))
                expected_type = column.get("type") or "string"
                if expected_type != detected_type and not (expected_type == "float" and detected_type == "integer"):
                    mismatches.append({"column": name, "expected": expected_type, "detected": detected_type})

        return self.result(
            passed=len(missing) == 0 and len(mismatches) == 0,
            severity=rule_config.get("severity", "blocking"),
            message="Le schema de sortie respecte les colonnes attendues."
            if len(missing) == 0 and len(mismatches) == 0
            else "Le fichier de sortie ne respecte pas totalement le schema attendu.",
            details={"missing_columns": missing, "type_mismatches": mismatches, "detected_headers": sorted(headers)},
        )
